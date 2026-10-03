#!/usr/bin/env python3

from copy import copy
from pathlib import Path
import csv
import shutil

from openpyxl import load_workbook
from openpyxl.styles.cell_style import StyleArray
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[2]

SOURCE_WORKBOOK = (
    ROOT
    / "supplementary_tables"
    / "Hsp60_10_Supplementary_Tables.xlsx"
)

OUTPUT_WORKBOOK = (
    ROOT
    / "supplementary_tables"
    / "Hsp60_10_Supplementary_Tables_REVISED.xlsx"
)

DEMOGRAPHICS_FILE = (
    ROOT
    / "outputs"
    / "reviewer_revisions"
    / "demographic_tables"
    / "demographics_manuscript_ready_table.csv"
)

EXPECTED_ORIGINAL_SHEETS = [
    "ST1_Client_inventory",
    "ST2_Pathway_coverage",
    "ST3_Fig5_null_input",
    "ST4_Fig5_null_summary",
    "ST5_Fig6_client_priority",
    "ST6_Fig6_layer_summary",
    "ST7_Sensitivity_checks",
]

ST8_NAME = "ST8_Demographics_APOE"
ST9_NAME = "ST9_Conventional_DE_DA"
ST10_NAME = "ST10_Pathology_models"
ST11_NAME = "ST11_Network_cognition"
ST12_NAME = "ST12_APOE_sensitivity"

APOE_EFFECT_SUMMARY_FILE = (
    ROOT
    / "outputs"
    / "reviewer_revisions"
    / "APOE_sensitivity"
    / "APOE_sensitivity_effect_summary.csv"
)

APOE_VULNERABILITY_SUMMARY_FILE = (
    ROOT
    / "outputs"
    / "reviewer_revisions"
    / "APOE_sensitivity"
    / "APOE_sensitivity_vulnerability_summary.csv"
)

NETWORK_COGNITION_FILE = (
    ROOT
    / "outputs"
    / "reviewer_revisions"
    / "reporting_requirements_resolution"
    / "06_network_cognition_reporting_complete.csv"
)

PATHOLOGY_REPORTING_FILE = (
    ROOT
    / "outputs"
    / "reviewer_revisions"
    / "reporting_requirements_resolution"
    / "02_pathology_Braak_CERAD_effect_SE_CI_n_P_FDR.csv"
)

CONVENTIONAL_DIR = (
    ROOT
    / "outputs"
    / "reviewer_revisions"
    / "conventional_stage_differential_final"
)

CLIENT_DIFFERENTIAL_FILE = (
    CONVENTIONAL_DIR
    / "14_PRIMARY_Hsp60_10_client_differential_results.csv"
)

HSPD1_HSPE1_FILE = (
    CONVENTIONAL_DIR
    / "15_PRIMARY_HSPD1_HSPE1_results.csv"
)

ST8_SOURCE_COLUMNS = [
    "row",
    "RNA_corrected_577__Overall",
    "RNA_corrected_577__NCI",
    "RNA_corrected_577__MCI",
    "RNA_corrected_577__AD",
    "Protein_stage_nuisance_complete_372__Overall",
    "Protein_stage_nuisance_complete_372__NCI",
    "Protein_stage_nuisance_complete_372__MCI",
    "Protein_stage_nuisance_complete_372__AD",
    "Regional_paired_DLPFC_STG_215__Overall",
]

ST8_DISPLAY_COLUMNS = [
    "Characteristic",
    "RNA overall",
    "RNA NCI",
    "RNA MCI",
    "RNA AD",
    "Protein overall",
    "Protein NCI",
    "Protein MCI",
    "Protein AD",
    "Regional paired cohort",
]

HEADER_FILL = "4F6B7A"
HEADER_FONT = "FFFFFF"
SECTION_FILL = "E9EEF1"
ALT_FILL = "F7F9FA"
BORDER_COLOR = "D4D9DD"

thin_border = Border(
    left=Side(style="thin", color=BORDER_COLOR),
    right=Side(style="thin", color=BORDER_COLOR),
    top=Side(style="thin", color=BORDER_COLOR),
    bottom=Side(style="thin", color=BORDER_COLOR),
)


def require_file(path):
    if not path.exists():
        raise FileNotFoundError(
            f"Required file does not exist:\n{path}"
        )


def load_demographics():
    with DEMOGRAPHICS_FILE.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as handle:
        rows = list(csv.DictReader(handle))

    if not rows:
        raise RuntimeError(
            "Demographics source table is empty."
        )

    available = set(rows[0].keys())

    missing = [
        col
        for col in ST8_SOURCE_COLUMNS
        if col not in available
    ]

    if missing:
        raise RuntimeError(
            "Demographics source is missing required "
            "column(s): "
            + ", ".join(missing)
        )

    expected_rows = [
        "N",
        "Age at death, mean (SD)",
        "Female, n (%)",
        "PMI, median [IQR]",
        "RIN, mean (SD)",
        "Braak, median [IQR]",
        "CERAD, median [IQR]",
        "APOE e4 carrier, n/N (%)",
        "APOE missing, n",
    ]

    observed_rows = [row["row"] for row in rows]

    missing_rows = [
        row
        for row in expected_rows
        if row not in observed_rows
    ]

    if missing_rows:
        raise RuntimeError(
            "Demographics source is missing required "
            "row(s): "
            + ", ".join(missing_rows)
        )

    row_lookup = {
        row["row"]: row
        for row in rows
    }

    return [
        row_lookup[name]
        for name in expected_rows
    ]


def snapshot_sheet(ws):
    """
    Record cell values/formulas for an audit that ST1-ST7
    were not changed while building the revised workbook.
    """
    return tuple(
        tuple(
            ws.cell(row=r, column=c).value
            for c in range(1, ws.max_column + 1)
        )
        for r in range(1, ws.max_row + 1)
    )


def style_st8(ws):
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    ws.sheet_view.showGridLines = False

    for cell in ws[1]:
        cell.fill = PatternFill(
            "solid",
            fgColor=HEADER_FILL,
        )
        cell.font = Font(
            bold=True,
            color=HEADER_FONT,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = thin_border

    ws.row_dimensions[1].height = 34

    for row_idx in range(
        2,
        ws.max_row + 1,
    ):
        fill = (
            PatternFill(
                "solid",
                fgColor=ALT_FILL,
            )
            if row_idx % 2 == 0
            else PatternFill(
                fill_type=None
            )
        )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            cell = ws.cell(
                row=row_idx,
                column=col_idx,
            )

            cell.fill = fill
            cell.border = thin_border
            cell.alignment = Alignment(
                vertical="center",
                horizontal=(
                    "left"
                    if col_idx == 1
                    else "center"
                ),
                wrap_text=True,
            )

        ws.cell(
            row=row_idx,
            column=1,
        ).font = Font(
            bold=True
        )

    widths = {
        1: 29,
        2: 16,
        3: 14,
        4: 14,
        5: 14,
        6: 17,
        7: 15,
        8: 15,
        9: 15,
        10: 20,
    }

    for col_idx, width in widths.items():
        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = width

    for row_idx in range(
        2,
        ws.max_row + 1,
    ):
        ws.row_dimensions[row_idx].height = 23


def build_st8(wb, demographics):
    if ST8_NAME in wb.sheetnames:
        del wb[ST8_NAME]

    ws = wb.create_sheet(ST8_NAME)

    for col_idx, value in enumerate(
        ST8_DISPLAY_COLUMNS,
        start=1,
    ):
        ws.cell(
            row=1,
            column=col_idx,
            value=value,
        )

    for row_idx, source_row in enumerate(
        demographics,
        start=2,
    ):
        for col_idx, source_col in enumerate(
            ST8_SOURCE_COLUMNS,
            start=1,
        ):
            value = source_row.get(
                source_col,
                "",
            )

            if value == "NA":
                value = ""

            ws.cell(
                row=row_idx,
                column=col_idx,
                value=value,
            )

    style_st8(ws)

    return ws


def validate_st8(ws):
    expected_n = {
        "RNA overall": "577",
        "RNA NCI": "200",
        "RNA MCI": "158",
        "RNA AD": "219",
        "Protein overall": "372",
        "Protein NCI": "167",
        "Protein MCI": "96",
        "Protein AD": "109",
        "Regional paired cohort": "215",
    }

    header = {
        ws.cell(
            row=1,
            column=c,
        ).value: c
        for c in range(
            1,
            ws.max_column + 1,
        )
    }

    characteristic_rows = {
        ws.cell(
            row=r,
            column=1,
        ).value: r
        for r in range(
            2,
            ws.max_row + 1,
        )
    }

    if "N" not in characteristic_rows:
        raise RuntimeError(
            "ST8 validation failed: N row missing."
        )

    n_row = characteristic_rows["N"]

    for column_name, expected in expected_n.items():
        col_idx = header[column_name]

        observed = ws.cell(
            row=n_row,
            column=col_idx,
        ).value

        if str(observed) != expected:
            raise RuntimeError(
                "ST8 validation failed for "
                f"{column_name}: expected "
                f"{expected}, observed {observed}."
            )

    required_characteristics = {
        "Age at death, mean (SD)",
        "Female, n (%)",
        "PMI, median [IQR]",
        "RIN, mean (SD)",
        "Braak, median [IQR]",
        "CERAD, median [IQR]",
        "APOE e4 carrier, n/N (%)",
        "APOE missing, n",
    }

    missing = required_characteristics.difference(
        characteristic_rows
    )

    if missing:
        raise RuntimeError(
            "ST8 validation failed; missing "
            "characteristic(s): "
            + ", ".join(sorted(missing))
        )


def read_csv_rows(path):
    with path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as handle:
        return list(csv.DictReader(handle))


def parse_excel_value(value):
    if value is None:
        return None

    value = str(value).strip()

    if value == "" or value.upper() == "NA":
        return None

    upper = value.upper()

    if upper == "TRUE":
        return True

    if upper == "FALSE":
        return False

    try:
        return int(value)
    except ValueError:
        pass

    try:
        return float(value)
    except ValueError:
        return value


def load_st9_rows():
    require_file(CLIENT_DIFFERENTIAL_FILE)
    require_file(HSPD1_HSPE1_FILE)

    client_rows = read_csv_rows(
        CLIENT_DIFFERENTIAL_FILE
    )

    hsp_rows = read_csv_rows(
        HSPD1_HSPE1_FILE
    )

    required = {
        "modality",
        "universe",
        "contrast",
        "feature_id",
        "ensembl_id",
        "gene_symbol",
        "mapping_status",
        "base_mean",
        "effect",
        "effect_label",
        "std_error",
        "conf_low",
        "conf_high",
        "statistic",
        "p_value",
        "fdr",
        "significant_fdr05",
    }

    for label, rows in [
        ("client differential", client_rows),
        ("HSPD1/HSPE1", hsp_rows),
    ]:
        if not rows:
            raise RuntimeError(
                f"{label} source table is empty."
            )

        missing = required.difference(
            rows[0].keys()
        )

        if missing:
            raise RuntimeError(
                f"{label} source missing column(s): "
                + ", ".join(sorted(missing))
            )

    out = []

    for row in client_rows:
        x = dict(row)
        x["target_type"] = "Hsp60/10 client"
        out.append(x)

    for row in hsp_rows:
        x = dict(row)
        x["target_type"] = "HSPD1/HSPE1"
        out.append(x)

    contrast_order = {
        "MCI_vs_NCI": 0,
        "AD_vs_MCI": 1,
        "AD_vs_NCI": 2,
    }

    modality_order = {
        "RNA": 0,
        "Protein": 1,
    }

    target_order = {
        "Hsp60/10 client": 0,
        "HSPD1/HSPE1": 1,
    }

    out.sort(
        key=lambda x: (
            target_order.get(
                x["target_type"],
                99,
            ),
            x["gene_symbol"],
            contrast_order.get(
                x["contrast"],
                99,
            ),
            modality_order.get(
                x["modality"],
                99,
            ),
        )
    )

    return out


def build_st9(wb, rows):
    if ST9_NAME in wb.sheetnames:
        del wb[ST9_NAME]

    ws = wb.create_sheet(ST9_NAME)

    columns = [
        ("Target type", "target_type"),
        ("Gene", "gene_symbol"),
        ("Modality", "modality"),
        ("Contrast", "contrast"),
        ("Universe", "universe"),
        ("Feature ID", "feature_id"),
        ("Ensembl ID", "ensembl_id"),
        ("Mapping status", "mapping_status"),
        ("Base mean", "base_mean"),
        ("Effect", "effect"),
        ("Effect definition", "effect_label"),
        ("SE", "std_error"),
        ("95% CI lower", "conf_low"),
        ("95% CI upper", "conf_high"),
        ("Statistic", "statistic"),
        ("P value", "p_value"),
        ("FDR", "fdr"),
        ("FDR < 0.05", "significant_fdr05"),
    ]

    for col_idx, (label, _) in enumerate(
        columns,
        start=1,
    ):
        ws.cell(
            row=1,
            column=col_idx,
            value=label,
        )

    contrast_display = {
        "MCI_vs_NCI": "MCI vs NCI",
        "AD_vs_MCI": "AD vs MCI",
        "AD_vs_NCI": "AD vs NCI",
    }

    for row_idx, source in enumerate(
        rows,
        start=2,
    ):
        for col_idx, (_, field) in enumerate(
            columns,
            start=1,
        ):
            value = source.get(field)

            if field == "contrast":
                value = contrast_display.get(
                    value,
                    value,
                )
            else:
                value = parse_excel_value(
                    value
                )

            if field == "significant_fdr05":
                if value is True:
                    value = "Yes"
                elif value is False:
                    value = "No"

            ws.cell(
                row=row_idx,
                column=col_idx,
                value=value,
            )

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85

    for cell in ws[1]:
        cell.fill = PatternFill(
            "solid",
            fgColor=HEADER_FILL,
        )
        cell.font = Font(
            bold=True,
            color=HEADER_FONT,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = thin_border

    ws.row_dimensions[1].height = 34

    widths = {
        1: 18,
        2: 14,
        3: 11,
        4: 14,
        5: 24,
        6: 20,
        7: 20,
        8: 19,
        9: 13,
        10: 13,
        11: 31,
        12: 13,
        13: 15,
        14: 15,
        15: 13,
        16: 14,
        17: 14,
        18: 14,
    }

    for col_idx, width in widths.items():
        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = width

    significant_fill = PatternFill(
        "solid",
        fgColor="FCE8E6",
    )

    significant_font = Font(
        bold=True,
        color="9C0006",
    )

    numeric_cols = {
        9, 10, 12, 13, 14, 15
    }

    pvalue_cols = {
        16, 17
    }

    for row_idx in range(
        2,
        ws.max_row + 1,
    ):
        alt = (
            row_idx % 2 == 0
        )

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            cell = ws.cell(
                row=row_idx,
                column=col_idx,
            )

            if alt:
                cell.fill = PatternFill(
                    "solid",
                    fgColor=ALT_FILL,
                )

            cell.border = thin_border

            cell.alignment = Alignment(
                vertical="center",
                horizontal=(
                    "left"
                    if col_idx in {
                        1, 2, 5, 6, 7, 8, 11
                    }
                    else "center"
                ),
                wrap_text=(
                    col_idx in {
                        1, 5, 6, 7, 8, 11
                    }
                ),
            )

            if col_idx in numeric_cols:
                cell.number_format = "0.0000"

            if col_idx in pvalue_cols:
                cell.number_format = "0.00E+00"

        sig_cell = ws.cell(
            row=row_idx,
            column=18,
        )

        if sig_cell.value == "Yes":
            sig_cell.fill = significant_fill
            sig_cell.font = significant_font

            ws.cell(
                row=row_idx,
                column=17,
            ).fill = significant_fill

            ws.cell(
                row=row_idx,
                column=17,
            ).font = significant_font

    return ws


def validate_st9(ws):
    expected_client_rows = 1713
    expected_hsp_rows = 12

    target_col = 1
    gene_col = 2
    modality_col = 3
    contrast_col = 4
    sig_col = 18

    rows = list(
        ws.iter_rows(
            min_row=2,
            values_only=True,
        )
    )

    client_rows = [
        row
        for row in rows
        if row[target_col - 1]
        == "Hsp60/10 client"
    ]

    hsp_rows = [
        row
        for row in rows
        if row[target_col - 1]
        == "HSPD1/HSPE1"
    ]

    if len(client_rows) != expected_client_rows:
        raise RuntimeError(
            "ST9 validation failed: expected "
            f"{expected_client_rows} client rows, "
            f"observed {len(client_rows)}."
        )

    if len(hsp_rows) != expected_hsp_rows:
        raise RuntimeError(
            "ST9 validation failed: expected "
            f"{expected_hsp_rows} HSPD1/HSPE1 rows, "
            f"observed {len(hsp_rows)}."
        )

    client_genes_by_modality = {}

    for modality in ["RNA", "Protein"]:
        genes = {
            row[gene_col - 1]
            for row in client_rows
            if row[modality_col - 1]
            == modality
        }

        client_genes_by_modality[
            modality
        ] = len(genes)

    expected_gene_counts = {
        "RNA": 297,
        "Protein": 274,
    }

    if (
        client_genes_by_modality
        != expected_gene_counts
    ):
        raise RuntimeError(
            "ST9 validation failed for modality "
            "gene counts. Expected "
            f"{expected_gene_counts}, observed "
            f"{client_genes_by_modality}."
        )

    hsp_genes = {
        row[gene_col - 1]
        for row in hsp_rows
    }

    if hsp_genes != {
        "HSPD1",
        "HSPE1",
    }:
        raise RuntimeError(
            "ST9 validation failed: expected "
            "HSPD1 and HSPE1 only; observed "
            f"{sorted(hsp_genes)}."
        )

    contrasts = {
        row[contrast_col - 1]
        for row in client_rows
    }

    if contrasts != {
        "MCI vs NCI",
        "AD vs MCI",
        "AD vs NCI",
    }:
        raise RuntimeError(
            "ST9 validation failed for contrasts: "
            f"{sorted(contrasts)}."
        )

    hspd1_ad_nci_protein = [
        row
        for row in hsp_rows
        if (
            row[gene_col - 1] == "HSPD1"
            and
            row[modality_col - 1]
            == "Protein"
            and
            row[contrast_col - 1]
            == "AD vs NCI"
        )
    ]

    if len(hspd1_ad_nci_protein) != 1:
        raise RuntimeError(
            "ST9 validation failed: expected one "
            "HSPD1 protein AD-vs-NCI row."
        )

    if (
        hspd1_ad_nci_protein[0][
            sig_col - 1
        ]
        != "Yes"
    ):
        raise RuntimeError(
            "ST9 validation failed: HSPD1 protein "
            "AD-vs-NCI should be FDR significant."
        )



def load_st10_rows():
    require_file(PATHOLOGY_REPORTING_FILE)

    rows = read_csv_rows(
        PATHOLOGY_REPORTING_FILE
    )

    if not rows:
        raise RuntimeError(
            "Pathology reporting source table is empty."
        )

    required = {
        "gene",
        "endpoint",
        "effect",
        "effect_label",
        "std_error",
        "conf_low_95",
        "conf_high_95",
        "n",
        "p_value",
        "fdr_bh",
    }

    missing = required.difference(
        rows[0].keys()
    )

    if missing:
        raise RuntimeError(
            "Pathology reporting source missing "
            "column(s): "
            + ", ".join(sorted(missing))
        )

    endpoint_order = {
        "Braak": 0,
        "CERAD": 1,
    }

    rows.sort(
        key=lambda x: (
            x["gene"],
            endpoint_order.get(
                x["endpoint"],
                99,
            ),
        )
    )

    return rows


def build_st10(wb, rows):
    if ST10_NAME in wb.sheetnames:
        del wb[ST10_NAME]

    ws = wb.create_sheet(ST10_NAME)

    columns = [
        ("Gene", "gene"),
        ("Endpoint", "endpoint"),
        ("Effect", "effect"),
        ("Effect definition", "effect_label"),
        ("SE", "std_error"),
        ("95% CI lower", "conf_low_95"),
        ("95% CI upper", "conf_high_95"),
        ("N", "n"),
        ("P value", "p_value"),
        ("BH FDR", "fdr_bh"),
    ]

    for col_idx, (label, _) in enumerate(
        columns,
        start=1,
    ):
        ws.cell(
            row=1,
            column=col_idx,
            value=label,
        )

    for row_idx, source in enumerate(
        rows,
        start=2,
    ):
        for col_idx, (_, field) in enumerate(
            columns,
            start=1,
        ):
            ws.cell(
                row=row_idx,
                column=col_idx,
                value=parse_excel_value(
                    source.get(field)
                ),
            )

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90

    for cell in ws[1]:
        cell.fill = PatternFill(
            "solid",
            fgColor=HEADER_FILL,
        )
        cell.font = Font(
            bold=True,
            color=HEADER_FONT,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = thin_border

    ws.row_dimensions[1].height = 34

    widths = {
        1: 15,
        2: 12,
        3: 14,
        4: 62,
        5: 14,
        6: 16,
        7: 16,
        8: 10,
        9: 14,
        10: 14,
    }

    for col_idx, width in widths.items():
        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = width

    sig_fill = PatternFill(
        "solid",
        fgColor="FCE8E6",
    )

    sig_font = Font(
        bold=True,
        color="9C0006",
    )

    for row_idx in range(
        2,
        ws.max_row + 1,
    ):
        alt = row_idx % 2 == 0

        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            cell = ws.cell(
                row=row_idx,
                column=col_idx,
            )

            if alt:
                cell.fill = PatternFill(
                    "solid",
                    fgColor=ALT_FILL,
                )

            cell.border = thin_border

            cell.alignment = Alignment(
                vertical="center",
                horizontal=(
                    "left"
                    if col_idx in {1, 4}
                    else "center"
                ),
                wrap_text=(col_idx == 4),
            )

        for col_idx in [3, 5, 6, 7]:
            ws.cell(
                row=row_idx,
                column=col_idx,
            ).number_format = "0.0000"

        for col_idx in [9, 10]:
            ws.cell(
                row=row_idx,
                column=col_idx,
            ).number_format = "0.00E+00"

        fdr_cell = ws.cell(
            row=row_idx,
            column=10,
        )

        if (
            isinstance(
                fdr_cell.value,
                (int, float),
            )
            and fdr_cell.value < 0.05
        ):
            fdr_cell.fill = sig_fill
            fdr_cell.font = sig_font

    return ws


def validate_st10(ws):
    rows = list(
        ws.iter_rows(
            min_row=2,
            values_only=True,
        )
    )

    if len(rows) != 612:
        raise RuntimeError(
            "ST10 validation failed: expected "
            f"612 model rows, observed {len(rows)}."
        )

    genes = {
        row[0]
        for row in rows
    }

    if len(genes) != 306:
        raise RuntimeError(
            "ST10 validation failed: expected "
            f"306 genes, observed {len(genes)}."
        )

    endpoints = {
        row[1]
        for row in rows
    }

    if endpoints != {
        "Braak",
        "CERAD",
    }:
        raise RuntimeError(
            "ST10 validation failed: unexpected "
            f"endpoints {sorted(endpoints)}."
        )

    counts = {}

    for endpoint in endpoints:
        counts[endpoint] = sum(
            row[1] == endpoint
            for row in rows
        )

    if counts != {
        "Braak": 306,
        "CERAD": 306,
    }:
        raise RuntimeError(
            "ST10 validation failed for endpoint "
            f"counts: {counts}."
        )

    gene_endpoint_counts = {}

    for row in rows:
        gene = row[0]
        endpoint = row[1]

        gene_endpoint_counts.setdefault(
            gene,
            set(),
        ).add(endpoint)

    bad_genes = [
        gene
        for gene, eps
        in gene_endpoint_counts.items()
        if eps != {
            "Braak",
            "CERAD",
        }
    ]

    if bad_genes:
        raise RuntimeError(
            "ST10 validation failed: genes without "
            "both Braak and CERAD results: "
            + ", ".join(
                sorted(bad_genes)[:20]
            )
        )


def load_st11_rows():
    require_file(NETWORK_COGNITION_FILE)

    rows = read_csv_rows(
        NETWORK_COGNITION_FILE
    )

    if not rows:
        raise RuntimeError(
            "Network cognition reporting table is empty."
        )

    required = {
        "outcome",
        "n",
        "estimate",
        "std.error",
        "statistic",
        "p.value",
        "conf.low",
        "conf.high",
        "r_squared",
        "adj_r_squared",
        "fdr_bh_three_network_outcomes",
        "multiplicity_note",
    }

    missing = required.difference(
        rows[0].keys()
    )

    if missing:
        raise RuntimeError(
            "Network cognition source missing "
            "column(s): "
            + ", ".join(sorted(missing))
        )

    outcome_order = {
        "Final cognitive diagnosis": 0,
        "Last-valid cognitive diagnosis": 1,
        "Last-valid MMSE": 2,
    }

    rows.sort(
        key=lambda x: outcome_order.get(
            x["outcome"],
            99,
        )
    )

    return rows


def build_st11(wb, rows):
    if ST11_NAME in wb.sheetnames:
        del wb[ST11_NAME]

    ws = wb.create_sheet(ST11_NAME)

    columns = [
        ("Outcome", "outcome"),
        ("N", "n"),
        ("Estimate", "estimate"),
        ("SE", "std.error"),
        ("Statistic", "statistic"),
        ("P value", "p.value"),
        ("95% CI lower", "conf.low"),
        ("95% CI upper", "conf.high"),
        ("R²", "r_squared"),
        ("Adjusted R²", "adj_r_squared"),
        (
            "BH FDR",
            "fdr_bh_three_network_outcomes",
        ),
        (
            "Multiplicity note",
            "multiplicity_note",
        ),
    ]

    for col_idx, (label, _) in enumerate(
        columns,
        start=1,
    ):
        ws.cell(
            row=1,
            column=col_idx,
            value=label,
        )

    for row_idx, source in enumerate(
        rows,
        start=2,
    ):
        for col_idx, (_, field) in enumerate(
            columns,
            start=1,
        ):
            ws.cell(
                row=row_idx,
                column=col_idx,
                value=parse_excel_value(
                    source.get(field)
                ),
            )

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90

    for cell in ws[1]:
        cell.fill = PatternFill(
            "solid",
            fgColor=HEADER_FILL,
        )
        cell.font = Font(
            bold=True,
            color=HEADER_FONT,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = thin_border

    ws.row_dimensions[1].height = 36

    widths = {
        1: 30,
        2: 10,
        3: 14,
        4: 14,
        5: 14,
        6: 14,
        7: 16,
        8: 16,
        9: 14,
        10: 16,
        11: 14,
        12: 65,
    }

    for col_idx, width in widths.items():
        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = width

    sig_fill = PatternFill(
        "solid",
        fgColor="FCE8E6",
    )

    sig_font = Font(
        bold=True,
        color="9C0006",
    )

    for row_idx in range(
        2,
        ws.max_row + 1,
    ):
        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            cell = ws.cell(
                row=row_idx,
                column=col_idx,
            )

            if row_idx % 2 == 0:
                cell.fill = PatternFill(
                    "solid",
                    fgColor=ALT_FILL,
                )

            cell.border = thin_border

            cell.alignment = Alignment(
                vertical="center",
                horizontal=(
                    "left"
                    if col_idx in {1, 12}
                    else "center"
                ),
                wrap_text=(col_idx in {1, 12}),
            )

        for col_idx in [
            3, 4, 5, 7, 8, 9, 10
        ]:
            ws.cell(
                row=row_idx,
                column=col_idx,
            ).number_format = "0.0000"

        for col_idx in [6, 11]:
            ws.cell(
                row=row_idx,
                column=col_idx,
            ).number_format = "0.00E+00"

        fdr_cell = ws.cell(
            row=row_idx,
            column=11,
        )

        if (
            isinstance(
                fdr_cell.value,
                (int, float),
            )
            and fdr_cell.value < 0.05
        ):
            fdr_cell.fill = sig_fill
            fdr_cell.font = sig_font

    return ws


def validate_st11(ws):
    rows = list(
        ws.iter_rows(
            min_row=2,
            values_only=True,
        )
    )

    if len(rows) != 3:
        raise RuntimeError(
            "ST11 validation failed: expected "
            f"3 network cognition models, "
            f"observed {len(rows)}."
        )

    expected_outcomes = {
        "Final cognitive diagnosis",
        "Last-valid cognitive diagnosis",
        "Last-valid MMSE",
    }

    observed_outcomes = {
        row[0]
        for row in rows
    }

    if observed_outcomes != expected_outcomes:
        raise RuntimeError(
            "ST11 validation failed for outcomes. "
            f"Observed: {sorted(observed_outcomes)}"
        )

    expected_n = {
        "Final cognitive diagnosis": 390,
        "Last-valid cognitive diagnosis": 391,
        "Last-valid MMSE": 398,
    }

    for row in rows:
        outcome = row[0]
        observed_n = row[1]

        if observed_n != expected_n[outcome]:
            raise RuntimeError(
                "ST11 validation failed for "
                f"{outcome}: expected n="
                f"{expected_n[outcome]}, "
                f"observed {observed_n}."
            )

        fdr = row[10]

        if not (
            isinstance(fdr, (int, float))
            and fdr < 0.05
        ):
            raise RuntimeError(
                "ST11 validation failed: "
                f"{outcome} should have "
                "BH FDR < 0.05."
            )


def load_st12_data():
    require_file(APOE_EFFECT_SUMMARY_FILE)
    require_file(APOE_VULNERABILITY_SUMMARY_FILE)

    effect_rows = read_csv_rows(
        APOE_EFFECT_SUMMARY_FILE
    )

    vulnerability_rows = read_csv_rows(
        APOE_VULNERABILITY_SUMMARY_FILE
    )

    if not effect_rows:
        raise RuntimeError(
            "APOE effect summary is empty."
        )

    if not vulnerability_rows:
        raise RuntimeError(
            "APOE vulnerability summary is empty."
        )

    required_effect = {
        "endpoint",
        "n_shared",
        "pearson_effect",
        "spearman_effect",
        "direction_concordance",
        "mean_effect_without_APOE",
        "mean_effect_with_APOE",
        "median_absolute_effect_change",
        "max_absolute_effect_change",
        "n_nominal_without_APOE",
        "n_nominal_with_APOE",
        "n_FDR_without_APOE",
        "n_FDR_with_APOE",
    }

    missing_effect = required_effect.difference(
        effect_rows[0].keys()
    )

    if missing_effect:
        raise RuntimeError(
            "APOE effect summary missing column(s): "
            + ", ".join(sorted(missing_effect))
        )

    required_vulnerability = {
        "metric",
        "spearman",
        "pearson",
    }

    missing_vulnerability = (
        required_vulnerability.difference(
            vulnerability_rows[0].keys()
        )
    )

    if missing_vulnerability:
        raise RuntimeError(
            "APOE vulnerability summary missing "
            "column(s): "
            + ", ".join(
                sorted(missing_vulnerability)
            )
        )

    return (
        effect_rows,
        vulnerability_rows,
    )


def build_st12(
    wb,
    effect_rows,
    vulnerability_rows,
):
    if ST12_NAME in wb.sheetnames:
        del wb[ST12_NAME]

    ws = wb.create_sheet(ST12_NAME)

    effect_columns = [
        ("Endpoint", "endpoint"),
        ("N shared", "n_shared"),
        ("Pearson r", "pearson_effect"),
        ("Spearman rho", "spearman_effect"),
        (
            "Direction concordance",
            "direction_concordance",
        ),
        (
            "Mean effect without APOE",
            "mean_effect_without_APOE",
        ),
        (
            "Mean effect with APOE",
            "mean_effect_with_APOE",
        ),
        (
            "Median |effect change|",
            "median_absolute_effect_change",
        ),
        (
            "Maximum |effect change|",
            "max_absolute_effect_change",
        ),
        (
            "Nominal P<0.05 without APOE",
            "n_nominal_without_APOE",
        ),
        (
            "Nominal P<0.05 with APOE",
            "n_nominal_with_APOE",
        ),
        (
            "FDR<0.05 without APOE",
            "n_FDR_without_APOE",
        ),
        (
            "FDR<0.05 with APOE",
            "n_FDR_with_APOE",
        ),
    ]

    ws.cell(
        row=1,
        column=1,
        value="A. APOE-adjustment sensitivity of client-level effects",
    )

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(effect_columns),
    )

    for col_idx, (label, _) in enumerate(
        effect_columns,
        start=1,
    ):
        ws.cell(
            row=2,
            column=col_idx,
            value=label,
        )

    for row_idx, source in enumerate(
        effect_rows,
        start=3,
    ):
        for col_idx, (_, field) in enumerate(
            effect_columns,
            start=1,
        ):
            value = parse_excel_value(
                source.get(field)
            )

            if field == "endpoint":
                value = {
                    "Late_AD_minus_MCI":
                        "Late AD minus MCI",
                }.get(
                    value,
                    value,
                )

            ws.cell(
                row=row_idx,
                column=col_idx,
                value=value,
            )

    vulnerability_start = (
        3
        + len(effect_rows)
        + 2
    )

    ws.cell(
        row=vulnerability_start,
        column=1,
        value=(
            "B. Concordance of derived vulnerability "
            "metrics with and without APOE adjustment"
        ),
    )

    ws.merge_cells(
        start_row=vulnerability_start,
        start_column=1,
        end_row=vulnerability_start,
        end_column=3,
    )

    vulnerability_header = (
        vulnerability_start + 1
    )

    vulnerability_columns = [
        ("Metric", "metric"),
        ("Spearman rho", "spearman"),
        ("Pearson r", "pearson"),
    ]

    for col_idx, (label, _) in enumerate(
        vulnerability_columns,
        start=1,
    ):
        ws.cell(
            row=vulnerability_header,
            column=col_idx,
            value=label,
        )

    metric_display = {
        "late_decline_percentile":
            "Late-decline percentile",
        "inverse_braak_percentile":
            "Inverse-Braak percentile",
        "pathology_vulnerability_score":
            "Pathology-vulnerability score",
    }

    for row_idx, source in enumerate(
        vulnerability_rows,
        start=vulnerability_header + 1,
    ):
        for col_idx, (_, field) in enumerate(
            vulnerability_columns,
            start=1,
        ):
            value = parse_excel_value(
                source.get(field)
            )

            if field == "metric":
                value = metric_display.get(
                    value,
                    value,
                )

            ws.cell(
                row=row_idx,
                column=col_idx,
                value=value,
            )

    ws.freeze_panes = "B3"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85

    section_rows = {
        1,
        vulnerability_start,
    }

    header_rows = {
        2,
        vulnerability_header,
    }

    for row_idx in section_rows:
        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            cell = ws.cell(
                row=row_idx,
                column=col_idx,
            )

            cell.fill = PatternFill(
                "solid",
                fgColor=SECTION_FILL,
            )
            cell.font = Font(
                bold=True,
            )
            cell.alignment = Alignment(
                vertical="center",
                horizontal="left",
                wrap_text=True,
            )

    for row_idx in header_rows:
        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            cell = ws.cell(
                row=row_idx,
                column=col_idx,
            )

            if cell.value is None:
                continue

            cell.fill = PatternFill(
                "solid",
                fgColor=HEADER_FILL,
            )
            cell.font = Font(
                bold=True,
                color=HEADER_FONT,
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = thin_border

    data_rows = (
        list(
            range(
                3,
                3 + len(effect_rows),
            )
        )
        + list(
            range(
                vulnerability_header + 1,
                vulnerability_header
                + 1
                + len(vulnerability_rows),
            )
        )
    )

    for row_idx in data_rows:
        for col_idx in range(
            1,
            ws.max_column + 1,
        ):
            cell = ws.cell(
                row=row_idx,
                column=col_idx,
            )

            if cell.value is None:
                continue

            cell.border = thin_border
            cell.alignment = Alignment(
                vertical="center",
                horizontal=(
                    "left"
                    if col_idx == 1
                    else "center"
                ),
                wrap_text=True,
            )

    for row_idx in range(
        3,
        3 + len(effect_rows),
    ):
        for col_idx in range(
            3,
            10,
        ):
            ws.cell(
                row=row_idx,
                column=col_idx,
            ).number_format = "0.0000"

    for row_idx in range(
        vulnerability_header + 1,
        vulnerability_header
        + 1
        + len(vulnerability_rows),
    ):
        for col_idx in [2, 3]:
            ws.cell(
                row=row_idx,
                column=col_idx,
            ).number_format = "0.0000"

    widths = {
        1: 29,
        2: 12,
        3: 13,
        4: 15,
        5: 20,
        6: 23,
        7: 21,
        8: 22,
        9: 22,
        10: 24,
        11: 22,
        12: 21,
        13: 19,
    }

    for col_idx, width in widths.items():
        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = width

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[
        vulnerability_start
    ].height = 32

    return ws


def validate_st12(
    ws,
    effect_rows,
    vulnerability_rows,
):
    if len(effect_rows) != 3:
        raise RuntimeError(
            "ST12 validation failed: expected "
            f"3 APOE effect-summary rows, "
            f"observed {len(effect_rows)}."
        )

    if len(vulnerability_rows) != 3:
        raise RuntimeError(
            "ST12 validation failed: expected "
            f"3 vulnerability-summary rows, "
            f"observed {len(vulnerability_rows)}."
        )

    endpoints = {
        row["endpoint"]
        for row in effect_rows
    }

    if endpoints != {
        "Late_AD_minus_MCI",
        "Braak",
        "CERAD",
    }:
        raise RuntimeError(
            "ST12 validation failed for APOE "
            f"effect endpoints: {sorted(endpoints)}."
        )

    for row in effect_rows:
        if int(row["n_shared"]) != 306:
            raise RuntimeError(
                "ST12 validation failed: expected "
                "n_shared = 306 for every endpoint."
            )

    late = next(
        row
        for row in effect_rows
        if row["endpoint"]
        == "Late_AD_minus_MCI"
    )

    if float(
        late["direction_concordance"]
    ) != 1.0:
        raise RuntimeError(
            "ST12 validation failed: late-stage "
            "direction concordance should equal 1."
        )

    metrics = {
        row["metric"]
        for row in vulnerability_rows
    }

    if metrics != {
        "late_decline_percentile",
        "inverse_braak_percentile",
        "pathology_vulnerability_score",
    }:
        raise RuntimeError(
            "ST12 validation failed for "
            "vulnerability metrics: "
            f"{sorted(metrics)}."
        )


def clear_manual_table_formatting(ws, cell_range):
    """
    Remove manual fills/borders from newly generated tables so that
    Excel's TableStyleMedium2 controls their appearance, matching
    the original ST1-ST7 workbook style.
    """
    for row in ws[cell_range]:
        for cell in row:
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()
            cell.font = Font(
                name="Calibri",
                size=11,
                bold=False,
                color="000000",
            )
            cell.alignment = Alignment(
                vertical="center",
                horizontal="general",
                wrap_text=True,
            )


def add_original_table_style(
    ws,
    ref,
    table_name,
):
    """
    Apply the same Excel table style used by the original
    supplementary tables: TableStyleMedium2 with banded rows.
    """
    for existing_name in list(ws.tables.keys()):
        del ws.tables[existing_name]

    ws.auto_filter.ref = None

    clear_manual_table_formatting(
        ws,
        ref,
    )

    start_cell, end_cell = ref.split(":")
    end_col = "".join(
        char
        for char in end_cell
        if char.isalpha()
    )

    header = ws[
        f"A1:{end_col}1"
    ][0]

    for cell in header:
        cell.font = Font(
            name="Calibri",
            size=12,
            bold=True,
            color="FFFFFF",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    table = Table(
        displayName=table_name,
        ref=ref,
    )

    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    ws.add_table(table)


def copy_cell_style(source, target):
    target._style = copy(source._style)
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def style_standard_supplementary_table(
    ws,
    reference_ws,
    header_row=1,
    data_start_row=2,
):
    """
    Copy the actual formatting used in ST1-ST7 rather than
    relying only on an Excel table-style name.
    """

    reference_header = reference_ws["A1"]
    reference_body = reference_ws["A2"]

    for cell in ws[header_row]:
        if cell.value is not None:
            copy_cell_style(
                reference_header,
                cell,
            )

    for row in ws.iter_rows(
        min_row=data_start_row,
        max_row=ws.max_row,
        min_col=1,
        max_col=ws.max_column,
    ):
        for cell in row:
            copy_cell_style(
                reference_body,
                cell,
            )

    ws.row_dimensions[
        header_row
    ].height = reference_ws.row_dimensions[1].height

    for row_idx in range(
        data_start_row,
        ws.max_row + 1,
    ):
        ws.row_dimensions[
            row_idx
        ].height = reference_ws.row_dimensions[2].height

    ws.sheet_view.showGridLines = (
        reference_ws.sheet_view.showGridLines
    )

    ws.freeze_panes = reference_ws.freeze_panes


def reset_direct_cell_formatting(
    ws,
    min_row=1,
    max_row=None,
    min_col=1,
    max_col=None,
):
    """
    Remove direct cell formatting so Excel's table style,
    rather than cell-level fills/fonts/borders, controls
    the visual appearance.

    This is necessary because ST1-ST7 use TableStyleMedium2
    and Excel renders that style using the workbook theme.
    """
    if max_row is None:
        max_row = ws.max_row

    if max_col is None:
        max_col = ws.max_column

    for row in ws.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            cell._style = StyleArray()
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


def replace_with_medium2_table(
    ws,
    ref,
    table_name,
    show_filter=False,
):
    """
    Create a clean TableStyleMedium2 table with no
    direct fill/font formatting overriding the style.
    """
    for existing_name in list(
        ws.tables.keys()
    ):
        del ws.tables[existing_name]

    ws.auto_filter.ref = None

    table = Table(
        displayName=table_name,
        ref=ref,
    )

    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    table.showAutoFilter = show_filter

    ws.add_table(table)


def apply_original_supplementary_style(wb):
    """
    Match ST8-ST12 to ST1-ST7 by allowing Excel's
    TableStyleMedium2 to control the actual colors.

    Do NOT directly paint the header cells.
    """

    ws8 = wb[ST8_NAME]
    reset_direct_cell_formatting(ws8)
    replace_with_medium2_table(
        ws8,
        f"A1:J{ws8.max_row}",
        "ST8DemographicsAPOE",
        show_filter=False,
    )

    ws9 = wb[ST9_NAME]
    reset_direct_cell_formatting(ws9)
    replace_with_medium2_table(
        ws9,
        f"A1:R{ws9.max_row}",
        "ST9ConventionalDEDA",
        show_filter=False,
    )

    ws10 = wb[ST10_NAME]
    reset_direct_cell_formatting(ws10)
    replace_with_medium2_table(
        ws10,
        f"A1:J{ws10.max_row}",
        "ST10PathologyModels",
        show_filter=False,
    )

    ws11 = wb[ST11_NAME]
    reset_direct_cell_formatting(ws11)
    replace_with_medium2_table(
        ws11,
        f"A1:L{ws11.max_row}",
        "ST11NetworkCognition",
        show_filter=False,
    )

    ws12 = wb[ST12_NAME]

    for existing_name in list(
        ws12.tables.keys()
    ):
        del ws12.tables[existing_name]

    reset_direct_cell_formatting(ws12)

    # ST12 contains two logical tables.
    effect_table = Table(
        displayName="ST12APOEEffects",
        ref="A2:M5",
    )

    effect_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    effect_table.showAutoFilter = False

    ws12.add_table(
        effect_table
    )

    vulnerability_table = Table(
        displayName="ST12Vulnerability",
        ref="A9:C12",
    )

    vulnerability_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    vulnerability_table.showAutoFilter = False

    ws12.add_table(
        vulnerability_table
    )

    # These are section labels OUTSIDE the Excel tables,
    # so a small amount of direct formatting is appropriate.
    for row_idx in [1, 8]:
        cell = ws12.cell(
            row=row_idx,
            column=1,
        )

        cell.font = Font(
            name="Calibri",
            size=11,
            bold=True,
            color="000000",
        )

        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

    # Restore useful number formats. Number formats do not
    # interfere with the Excel table colors.
    for row in range(
        2,
        ws9.max_row + 1,
    ):
        for col in range(
            9,
            16,
        ):
            ws9.cell(
                row=row,
                column=col,
            ).number_format = "0.0000"

        for col in [16, 17]:
            ws9.cell(
                row=row,
                column=col,
            ).number_format = "0.00E+00"

    for row in range(
        2,
        ws10.max_row + 1,
    ):
        for col in range(
            3,
            8,
        ):
            ws10.cell(
                row=row,
                column=col,
            ).number_format = "0.0000"

        for col in [9, 10]:
            ws10.cell(
                row=row,
                column=col,
            ).number_format = "0.00E+00"

    for row in range(
        2,
        ws11.max_row + 1,
    ):
        for col in [
            3,
            4,
            5,
            7,
            8,
            9,
            10,
        ]:
            ws11.cell(
                row=row,
                column=col,
            ).number_format = "0.0000"

        for col in [6, 11]:
            ws11.cell(
                row=row,
                column=col,
            ).number_format = "0.00E+00"

    for row in range(
        3,
        6,
    ):
        for col in range(
            3,
            10,
        ):
            ws12.cell(
                row=row,
                column=col,
            ).number_format = "0.0000"

    for row in range(
        10,
        13,
    ):
        for col in [2, 3]:
            ws12.cell(
                row=row,
                column=col,
            ).number_format = "0.0000"


def main():
    require_file(SOURCE_WORKBOOK)
    require_file(DEMOGRAPHICS_FILE)
    require_file(CLIENT_DIFFERENTIAL_FILE)
    require_file(HSPD1_HSPE1_FILE)
    require_file(PATHOLOGY_REPORTING_FILE)
    require_file(NETWORK_COGNITION_FILE)
    require_file(APOE_EFFECT_SUMMARY_FILE)
    require_file(APOE_VULNERABILITY_SUMMARY_FILE)

    demographics = load_demographics()
    st9_rows = load_st9_rows()
    st10_rows = load_st10_rows()
    st11_rows = load_st11_rows()
    (
        st12_effect_rows,
        st12_vulnerability_rows,
    ) = load_st12_data()

    wb_source = load_workbook(
        SOURCE_WORKBOOK,
        data_only=False,
    )

    if wb_source.sheetnames != EXPECTED_ORIGINAL_SHEETS:
        raise RuntimeError(
            "Canonical supplemental workbook sheets "
            "do not match the expected ST1-ST7 layout.\n"
            f"Expected: {EXPECTED_ORIGINAL_SHEETS}\n"
            f"Observed: {wb_source.sheetnames}"
        )

    source_snapshots = {
        name: snapshot_sheet(
            wb_source[name]
        )
        for name in EXPECTED_ORIGINAL_SHEETS
    }

    wb_source.close()

    shutil.copy2(
        SOURCE_WORKBOOK,
        OUTPUT_WORKBOOK,
    )

    wb = load_workbook(
        OUTPUT_WORKBOOK,
        data_only=False,
    )

    build_st8(
        wb,
        demographics,
    )

    validate_st8(
        wb[ST8_NAME]
    )

    build_st9(
        wb,
        st9_rows,
    )

    validate_st9(
        wb[ST9_NAME]
    )

    build_st10(
        wb,
        st10_rows,
    )

    validate_st10(
        wb[ST10_NAME]
    )

    build_st11(
        wb,
        st11_rows,
    )

    validate_st11(
        wb[ST11_NAME]
    )

    build_st12(
        wb,
        st12_effect_rows,
        st12_vulnerability_rows,
    )

    validate_st12(
        wb[ST12_NAME],
        st12_effect_rows,
        st12_vulnerability_rows,
    )

    for sheet_name in EXPECTED_ORIGINAL_SHEETS:
        if snapshot_sheet(
            wb[sheet_name]
        ) != source_snapshots[sheet_name]:
            raise RuntimeError(
                f"{sheet_name} changed unexpectedly "
                "while building the revised workbook."
            )

    apply_original_supplementary_style(
        wb
    )

    expected_final_sheets = (
        EXPECTED_ORIGINAL_SHEETS
        + [
            ST8_NAME,
            ST9_NAME,
            ST10_NAME,
            ST11_NAME,
            ST12_NAME,
        ]
    )

    if wb.sheetnames != expected_final_sheets:
        raise RuntimeError(
            "Unexpected final workbook sheet order.\n"
            f"Expected: {expected_final_sheets}\n"
            f"Observed: {wb.sheetnames}"
        )

    wb.save(
        OUTPUT_WORKBOOK
    )
    wb.close()

    print(
        "Supplementary table workbook built successfully."
    )
    print(
        f"Source: {SOURCE_WORKBOOK}"
    )
    print(
        f"ST8 source: {DEMOGRAPHICS_FILE}"
    )
    print(
        f"Output: {OUTPUT_WORKBOOK}"
    )
    print(
        "Verified: ST1-ST7 cell contents unchanged."
    )
    print(
        "Verified: ST8 cohort counts and required "
        "demographic rows."
    )
    print(
        "Verified: ST9 contains 1,713 client tests "
        "(297 RNA / 274 protein clients across "
        "three contrasts) plus 12 HSPD1/HSPE1 tests."
    )
    print(
        "Verified: ST10 contains 612 pathology "
        "models: 306 Braak and 306 CERAD results."
    )
    print(
        "Verified: ST11 contains all 3 "
        "prespecified network-level cognition "
        "models with n, estimate, SE, CI, P, "
        "R-squared, adjusted R-squared, and FDR."
    )
    print(
        "Verified: ST12 contains 3 APOE effect "
        "sensitivity summaries and 3 vulnerability "
        "metric concordance summaries."
    )
    print(
        "Verified: ST8-ST12 use the original "
        "TableStyleMedium2 formatting used by ST1-ST7."
    )


if __name__ == "__main__":
    main()
